"""Tests for model_change_impact.requirement_report - the four requirement
traceability sheets added to the Baseline Estimation workbook."""
from openpyxl import Workbook

from model_change_impact import requirement_report

RUN_AT = "2026-09-14T12:00:00"


def _requirements():
    return [
        {"id": "R-1", "title": "Revenue restatement", "status": "Approved", "active": True,
         "priority": "High", "business_area": "Finance"},
        {"id": "R-2", "title": "Old cleanup", "status": "Done", "active": False,
         "priority": "Low", "business_area": ""},
    ]


def _mapping():
    return {
        "R-1": {
            "objects": {
                ("measure", "Sales", "Total Revenue"): {"source": "Tag", "confidence": "High"},
                ("column", "Sales", "Amount"): {"source": "Inferred", "confidence": "Low"},
            },
            "visuals": {
                ("p1", "v1"): {"source": "Tag", "confidence": "High",
                               "page_display_name": "Overview",
                               "visual_display_name": "Revenue Card",
                               "visual_type": "card", "kpi_classification": "certain"},
            },
            "warnings": [],
        },
        "R-2": {"objects": {}, "visuals": {}, "warnings": []},
    }


def _history():
    return {
        "measure|Sales|Total Revenue": {
            "object_type": "Measure", "object_name": "Sales[Total Revenue]",
            "parent_object": "Sales", "first_seen": "2026-09-01T10:00:00",
            "last_changed": "2026-09-10T10:00:00",
            "previous_change_at": "2026-09-01T10:00:00", "change_count": 2,
            "last_change_type": "Modified",
        },
        "column|Sales|Amount": {
            "object_type": "Column", "object_name": "Sales[Amount]",
            "parent_object": "Sales", "first_seen": "2026-09-01T10:00:00",
            "last_changed": "2026-09-01T10:00:00", "previous_change_at": "",
            "change_count": 1, "last_change_type": "Added",
        },
        "visual|p1|v1": {
            "object_type": "Visual", "object_name": "Revenue Card",
            "parent_object": "Overview", "first_seen": "2026-09-01T10:00:00",
            "last_changed": "2026-09-01T10:00:00", "previous_change_at": "",
            "change_count": 1, "last_change_type": "Added",
        },
        "page|p1": {
            "object_type": "Page", "object_name": "Overview", "parent_object": "",
            "first_seen": "2026-09-01T10:00:00", "last_changed": "2026-09-01T10:00:00",
            "previous_change_at": "", "change_count": 1, "last_change_type": "Added",
        },
    }


def _events():
    return [
        {"object_key": "table|Sales", "object_type": "Table", "object_name": "Sales",
         "parent_object": "", "change_type": "Modified", "changed_at": "2026-09-13T09:00:00",
         "previous_change_at": "2026-09-01T09:00:00", "attributed_requirement_id": "Unmapped",
         "first_seen": "2026-09-01T09:00:00", "change_count": 3, "run_id": 2},
        {"object_key": "measure|Sales|Total Revenue", "object_type": "Measure",
         "object_name": "Sales[Total Revenue]", "parent_object": "Sales",
         "change_type": "Modified", "changed_at": "2026-09-10T10:00:00",
         "previous_change_at": "2026-09-01T10:00:00", "attributed_requirement_id": "R-1",
         "first_seen": "2026-09-01T10:00:00", "change_count": 2, "run_id": 2},
        {"object_key": "column|Sales|Amount", "object_type": "Column",
         "object_name": "Sales[Amount]", "parent_object": "Sales",
         "change_type": "Added", "changed_at": "2026-09-01T09:00:00",
         "previous_change_at": "", "attributed_requirement_id": "Unmapped",
         "first_seen": "2026-09-01T09:00:00", "change_count": 1, "run_id": 1},
    ]


def _build():
    workbook = Workbook()
    workbook.active.title = "Object Summary"
    stats = requirement_report.build_requirement_sheets(
        workbook, _requirements(), _mapping(), _history(), _events(),
        changed_keys={"measure|Sales|Total Revenue"}, run_at=RUN_AT)
    return workbook, stats


def _sheet_rows(workbook, sheet_name):
    return [tuple(row) for row in workbook[sheet_name].iter_rows(values_only=True)]


def test_four_sheets_created():
    workbook, _stats = _build()
    for name in ("Requirement Summary", "Requirement Traceability",
                 "Object Change History", "Unmapped Changes"):
        assert name in workbook.sheetnames


def test_requirement_summary_counts_and_severity():
    workbook, _stats = _build()
    rows = _sheet_rows(workbook, "Requirement Summary")
    header = rows[0]
    r1 = dict(zip(header, next(row for row in rows[1:] if row[0] == "R-1")))
    assert r1["Impacted Measures"] == 1
    assert r1["Impacted Columns"] == 1
    assert r1["Impacted Visuals"] == 1
    assert r1["Impacted Pages"] == 1
    assert r1["Total Impacted Objects"] == 4
    assert r1["Low Confidence Count"] == 1
    assert r1["Impact Severity"] == 4 * 3  # total x High weight
    assert r1["Business Area"] == "Finance"
    assert r1["Last Analyzed"] == RUN_AT
    r2 = dict(zip(header, next(row for row in rows[1:] if row[0] == "R-2")))
    assert r2["Total Impacted Objects"] == 0
    assert r2["Impact Severity"] == 0


def test_traceability_rows_and_review_flag():
    workbook, stats = _build()
    rows = _sheet_rows(workbook, "Requirement Traceability")
    header = rows[0]
    records = [dict(zip(header, row)) for row in rows[1:]]
    r1_records = [r for r in records if r["Requirement ID"] == "R-1"]

    measure_row = next(r for r in r1_records if r["Object Type"] == "Measure")
    assert measure_row["Object Name"] == "Sales[Total Revenue]"
    assert measure_row["Parent Object"] == "Sales"
    assert measure_row["Mapping Source"] == "Tag"
    assert measure_row["Change Status"] == "Modified"  # changed in this run
    assert measure_row["Days Since Last Change"] == 4  # Sep 10 -> Sep 14
    assert measure_row["Review Flag"] == ""

    column_row = next(r for r in r1_records if r["Object Type"] == "Column")
    assert column_row["Confidence"] == "Low"
    assert column_row["Review Flag"] == "Yes"
    assert column_row["Change Status"] == "Unchanged"  # recorded, not changed this run

    visual_row = next(r for r in r1_records if r["Object Type"] == "Visual")
    assert visual_row["Object Name"] == "Revenue Card"
    assert visual_row["Report Page"] == "Overview"

    page_row = next(r for r in r1_records if r["Object Type"] == "Page")
    assert page_row["Object Name"] == "Overview"
    assert page_row["Mapping Source"] == "Tag"  # inherits the tagged visual's source

    # Requirement with no mapped objects still gets a visible placeholder row.
    r2_row = next(r for r in records if r["Requirement ID"] == "R-2")
    assert r2_row["Object Type"] == "(no objects mapped)"

    assert stats["low_confidence"] == 1
    assert stats["mapped_objects"] == 2


def test_history_sheet_rows():
    workbook, _stats = _build()
    rows = _sheet_rows(workbook, "Object Change History")
    header = rows[0]
    assert len(rows) == 4  # header + 3 events
    first = dict(zip(header, rows[1]))
    assert first["Object Name"] == "Sales"
    assert first["Change Type"] == "Modified"
    assert first["Days Between Changes"] == 12  # Sep 1 -> Sep 13
    assert first["Attributed Requirement ID"] == "Unmapped"
    assert first["Change Count"] == 3


def test_unmapped_changes_sheet():
    workbook, stats = _build()
    rows = _sheet_rows(workbook, "Unmapped Changes")
    assert rows[0] == ("Object Type", "Object Name", "Change Type", "Changed At",
                       "Assign Requirement ID")
    assert len(rows) == 2  # header + 1 unmapped post-baseline event
    assert rows[1][0] == "Table" and rows[1][1] == "Sales"
    assert rows[1][4] in (None, "")
    assert stats["unmapped_changes"] == 1


def test_baseline_run_events_are_not_reviewable_changes():
    # The run_id=1 "Added" event (baseline registration) must not appear in
    # Unmapped Changes even though it is unattributed.
    _workbook, stats = _build()
    assert stats["unmapped_changes"] == 1
    events = _events()
    baseline_only = [event for event in events if event["run_id"] == 1]
    assert requirement_report._build_unmapped_rows(baseline_only) == []
