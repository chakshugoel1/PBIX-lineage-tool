"""Offline workflow tests: replace PBIX extraction, keep orchestration and Excel writing real."""
from copy import deepcopy
import json

from openpyxl import Workbook
from openpyxl import load_workbook

from gui import worker


def _snapshot(source_file):
    return {
        "source_file": source_file,
        "tables": {
            "Sales": {
                "is_calculated_table": False, "m_expression": "let Source = 1 in Source",
                "is_hidden": False, "description": None, "lineage_tag": None,
                "columns": [{
                    "name": "Amount", "data_type": "Int64", "is_calculated": False,
                    "expression": None, "format_string": None, "is_hidden": False,
                    "description": None, "display_folder": None, "lineage_tag": None,
                }],
            },
        },
        "measures": {"Sales": {"Total": {
            "expression": "SUM(Sales[Amount])", "display_folder": None,
            "description": None, "format_string": None, "is_hidden": False,
            "lineage_tag": None, "kpi_id": None,
        }}},
        "calculated_columns": [],
        "relationships": [],
    }


def _layout():
    return {
        "format": "pbir",
        "pages": [{
            "page_id": "p1", "display_name": "Overview",
            "visuals": [{
                "visual_id": "card1", "kind": "visual", "visual_type": "card",
                "display_name": "Total", "kpi_classification": "certain",
                "fields": [{"kind": "measure", "table": "Sales", "field": "Total", "role": "Values"}],
            }],
        }],
    }


def test_baseline_worker_writes_rtm_workbook(monkeypatch, tmp_path):
    monkeypatch.setattr(worker.snapshot, "build_snapshot", _snapshot)
    monkeypatch.setattr(worker.report_layout, "build_report_layout", lambda path: _layout())
    output_path = tmp_path / "baseline.xlsx"
    completed = []
    failures = []
    job = worker.BaselineEstimationWorker("baseline.pbix", str(output_path))
    job.finished_ok.connect(completed.append)
    job.failed.connect(failures.append)

    job.run()

    assert not failures
    assert completed and completed[0]["output_path"] == str(output_path)
    workbook = load_workbook(output_path)
    assert workbook.sheetnames == ["RTM", "Requirement Copy", "Visual Impact by Requirement"]
    rows = list(workbook["RTM"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][2:5] == ("card1", "Overview", "Total")
    assert (tmp_path / "visual_descriptions.json").exists()


def test_baseline_worker_maps_requirements(monkeypatch, tmp_path):
    monkeypatch.setattr(worker.snapshot, "build_snapshot", _snapshot)
    monkeypatch.setattr(worker.report_layout, "build_report_layout", lambda path: _layout())
    requirements_path = tmp_path / "requirements.xlsx"
    requirements = Workbook()
    sheet = requirements.active
    sheet.title = "Requirements"
    sheet.append(["Requirement ID", "Title", "Status", "Impacted Visual IDs"])
    sheet.append(["R-1", "Update total", "Approved", "card1"])
    requirements.save(requirements_path)
    output_path = tmp_path / "baseline.xlsx"
    completed = []
    failures = []
    job = worker.BaselineEstimationWorker(
        "baseline.pbix", str(output_path), requirements_path=str(requirements_path))
    job.finished_ok.connect(completed.append)
    job.failed.connect(failures.append)

    job.run()

    assert not failures
    assert completed and completed[0]["requirements"] == 1
    workbook = load_workbook(output_path)
    assert workbook["Requirement Copy"]["A2"].value == "R-1"
    assert workbook["Visual Impact by Requirement"]["F2"].value == "card1"


def test_model_change_worker_writes_impact_workbook(monkeypatch, tmp_path):
    def extract(path):
        model = _snapshot(path)
        if path == "changed.pbix":
            model = deepcopy(model)
            model["measures"]["Sales"]["Total"]["expression"] = "SUM(Sales[Amount]) * 2"
        return model

    monkeypatch.setattr(worker.snapshot, "build_snapshot", extract)
    monkeypatch.setattr(worker.report_layout, "build_report_layout", lambda path: _layout())
    output_path = tmp_path / "impact.xlsx"
    completed = []
    failures = []
    job = worker.ModelChangeImpactWorker("baseline.pbix", "changed.pbix", str(output_path))
    job.finished_ok.connect(completed.append)
    job.failed.connect(failures.append)

    job.run()

    assert not failures
    assert completed and completed[0]["measures"]["changed"] == 1
    assert completed[0]["impacted_visuals"] == 1
    workbook = load_workbook(output_path)
    assert workbook.sheetnames == [
        "Summary", "Impact Summary", "Changed Tables", "Changed Measures",
        "Changed Columns", "Changed Relationships",
    ]
    rows = list(workbook["Impact Summary"].iter_rows(min_row=2, values_only=True))
    assert any(row[0] == "Measure" and row[3] == "card1" and row[9] == "Direct" for row in rows)


def test_lineage_worker_writes_both_workbooks(monkeypatch, tmp_path):
    def extract(pbix_path, dataflow_folder, **kwargs):
        context = {
            "pbix_path": pbix_path, "entries": {}, "dataflows": {}, "report_layout": {},
            "model": type("Model", (), {"tables": ["Static"]})(),
        }
        row = worker.blr.resolve_table_row("Static", context)
        row.update({"table": "Static", "is_used": True})
        return [row], context

    monkeypatch.setattr(worker.blr, "build_report", extract)
    completed = []
    failures = []
    job = worker.PipelineWorker("sample.pbix", str(tmp_path), str(tmp_path), archive_previous=False)
    job.finished_ok.connect(completed.append)
    job.failed.connect(failures.append)

    job.run()

    assert not failures
    assert completed and completed[0]["total"] == 1
    main = load_workbook(completed[0]["generated_path"])
    companion = load_workbook(completed[0]["dataflow_lineage_path"])
    assert main.active["B2"].value == "No M/Power Query Source"
    assert companion["Table Lineage"]["A2"].value == "Static"
    assert companion["Table Lineage"]["C2"].value == "Calculated"


def test_export_worker_receives_offline_powershell_result(monkeypatch, tmp_path):
    output_path = tmp_path / "Sample.json"

    class FakeProcess:
        def __init__(self):
            output_path.write_text('{"name": "Sample"}', encoding="utf-8")
            result = {"success": True, "stage": "done", "message": "1 exported",
                      "files": [str(output_path)]}
            self.stdout = ["##RESULT##" + json.dumps(result)]

        def wait(self):
            return 0

    monkeypatch.setattr(worker.dataflow_export.subprocess, "Popen", lambda *args, **kwargs: FakeProcess())
    completed = []
    failures = []
    job = worker.DataflowExportWorker("workspace-1", str(tmp_path), archive_previous=False)
    job.finished_ok.connect(completed.append)
    job.failed.connect(failures.append)

    job.run()

    assert not failures
    assert completed and completed[0]["files"] == [str(output_path)]
    assert json.loads(output_path.read_text(encoding="utf-8")) == {"name": "Sample"}