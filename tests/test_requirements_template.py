from pathlib import Path

from openpyxl import load_workbook

from model_change_impact.requirements import load_requirements


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "Requirements_Template.xlsx"


def test_requirements_template_matches_loader_contract():
    workbook = load_workbook(TEMPLATE, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Requirements", "Instructions"]
    headers = [cell.value for cell in next(workbook["Requirements"].iter_rows())]
    assert headers[:4] == ["Requirement ID", "Requirement Type", "Title", "Status"]
    workbook.close()

    requirements, warnings = load_requirements(TEMPLATE)

    assert not warnings
    assert len(requirements) == 1
    assert requirements[0]["id"] == "REQ-EXAMPLE-001"
    assert requirements[0]["status"] == "Proposed"
    assert requirements[0]["priority"] == "Medium"