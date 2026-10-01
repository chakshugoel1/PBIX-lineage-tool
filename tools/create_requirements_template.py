"""Create the user-facing requirements workbook template."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "templates" / "Requirements_Template.xlsx"
HEADERS = [
    "Requirement ID", "Requirement Type", "Title", "Status", "Description",
    "Priority", "Raised Date", "Business Owner", "Business Area",
    "Expected Change Date", "Impacted Pages", "Impacted Visuals",
    "Impacted Visual IDs", "Notes",
]


def build_template(path=OUTPUT):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Requirements"
    sheet.append(HEADERS)
    sheet.append([
        "REQ-EXAMPLE-001", "Enhancement", "Replace this example requirement",
        "Proposed", "Delete or replace this example row before use.", "Medium",
        "2026-01-01", "", "", "", "Overview", "", "", "",
    ])
    sheet.append([None] * len(HEADERS))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:N200"
    sheet.row_dimensions[1].height = 30

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [20, 20, 34, 16, 48, 14, 15, 22, 20, 20, 24, 28, 28, 42]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2, max_row=200):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in ("G", "J"):
        for row in range(2, 201):
            sheet[f"{column}{row}"].number_format = "yyyy-mm-dd"

    validations = [
        ("B2:B200", '"New Requirement,Enhancement,Bug Fix"'),
        ("D2:D200", '"Proposed,Approved,In Progress,Done,Cancelled"'),
        ("F2:F200", '"Critical,High,Medium,Low"'),
    ]
    for cell_range, formula in validations:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(cell_range)

    comments = {
        "A1": "Required. Must be unique.",
        "D1": "Required. Use the dropdown values.",
        "F1": "Recommended. Defaults to Medium when blank.",
        "G1": "Recommended. Use YYYY-MM-DD.",
        "K1": "Optional. Use exact page names, separated by commas.",
        "L1": "Optional. Use exact visual names, separated by commas.",
        "M1": "Optional and preferred. Use exact Visual IDs from the RTM sheet.",
    }
    for coordinate, text in comments.items():
        sheet[coordinate].comment = Comment(text, "PBIX Lineage Tool")

    instructions = workbook.create_sheet("Instructions")
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 100
    instructions.append(["Topic", "Guidance"])
    instructions.append(["Purpose", "Use this workbook with the Baseline Estimation tab to map business requirements to report pages and visuals."])
    instructions.append(["Scope", "Prefer Impacted Visual IDs. Find them in the RTM sheet from an initial Baseline Estimation run."])
    instructions.append(["Mandatory", "Requirement ID, Title, and Status must be filled for every row."])
    instructions.append(["Dates", "Use YYYY-MM-DD, for example 2026-09-25."])
    instructions.append(["Model objects", "Do not enter tables, columns, or measures as scope. The tool derives them from visual bindings."])
    for cell in instructions[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for row in instructions.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    instructions.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


if __name__ == "__main__":
    print(build_template())